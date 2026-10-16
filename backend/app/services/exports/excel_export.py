import io
from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import Branch
from app.services.exports.csv_export import CsvExportService, sanitize_formula_injection


class ExcelExportService:
    @staticmethod
    def generate_xlsx_bytes(branches: List[Branch]) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "India Business Leads"

        # Styles
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        border_thin = Side(border_style="thin", color="E2E8F0")
        cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

        # Write header
        ws.append(CsvExportService.HEADERS)
        for col_idx in range(1, len(CsvExportService.HEADERS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

        # Write rows
        for row_idx, branch in enumerate(branches, start=2):
            row_data = CsvExportService.branch_to_row(branch)
            ws.append(row_data)

            # Style data cells
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = regular_font
                cell.border = cell_border
                cell.alignment = Alignment(vertical="center")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        # Freeze header row
        ws.freeze_panes = "A2"

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

