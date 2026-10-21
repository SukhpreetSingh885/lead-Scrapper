import io
import zipfile
import openpyxl
import pytest
from app.models import Branch, Company, BusinessContact, Export, Organization, DiscoveryJobItem
from app.workers.tasks.export_tasks import category_to_slug
from app.services.exports.csv_export import sanitize_formula_injection


def test_category_to_slug():
    assert category_to_slug("Hospital") == "hospital"
    assert category_to_slug("Dental Clinic") == "dental_clinic"
    assert category_to_slug("Pet Shop") == "pet_shop"
    assert category_to_slug("Fashion") == "fashion"
    assert category_to_slug("Salon") == "salon"
    assert category_to_slug("Gym & Fitness") == "gym_fitness"
    assert category_to_slug("") == "all"
    assert category_to_slug(None) == "all"


def test_formula_injection_defense():
    dangerous_inputs = [
        "=1+1",
        "+cmd|' /C calc'!A0",
        "-5+2",
        "@SUM(A1:A10)",
        "\t=evil()",
        "\r=evil()",
    ]
    for val in dangerous_inputs:
        sanitized = sanitize_formula_injection(val)
        assert sanitized.startswith("'"), f"Dangerous input {val} was not escaped: {sanitized}"

    # Normal values not altered
    assert sanitize_formula_injection("Apollo Hospital") == "Apollo Hospital"
    assert sanitize_formula_injection(None) == "Not Available"


def test_five_contact_types_qualification_matrix(client, db_session):
    """
    Rule 14: Test mobile-only, landline-only, WhatsApp-only, email-only, and no-contact cases.
    Verifies that:
      - Mobile only -> QUALIFIED
      - Landline only -> QUALIFIED
      - WhatsApp only -> QUALIFIED
      - Email only -> EXCLUDED
      - No contact -> EXCLUDED
      - Empty category -> Returns clear error, prevents empty file generation
    """
    org = db_session.query(Organization).first()
    if not org:
        org = Organization(id="org-test-qual", name="Qual Org", slug="qual-org")
        db_session.add(org)
        db_session.flush()

    comp = Company(
        organization_id=org.id,
        legal_name="Matrix Test Company",
        normalized_name="matrix test company",
        primary_category="Hospital"
    )
    db_session.add(comp)
    db_session.flush()

    # 1. Mobile-only lead
    b_mob = Branch(
        company_id=comp.id,
        branch_name="Branch Mobile Only",
        normalized_branch_name="branch mobile only",
        city="Mohali",
        state="Punjab",
        source_permission_status="permitted_storage"
    )
    db_session.add(b_mob)
    db_session.flush()
    db_session.add(BusinessContact(
        branch_id=b_mob.id,
        contact_type="mobile",
        original_value="+91 98765 11111",
        normalized_value="+919876511111"
    ))

    # 2. Landline-only lead
    b_land = Branch(
        company_id=comp.id,
        branch_name="Branch Landline Only",
        normalized_branch_name="branch landline only",
        city="Mohali",
        state="Punjab",
        source_permission_status="permitted_storage"
    )
    db_session.add(b_land)
    db_session.flush()
    db_session.add(BusinessContact(
        branch_id=b_land.id,
        contact_type="landline",
        original_value="0172 2701111",
        normalized_value="+911722701111"
    ))

    # 3. WhatsApp-only lead
    b_wa = Branch(
        company_id=comp.id,
        branch_name="Branch WhatsApp Only",
        normalized_branch_name="branch whatsapp only",
        city="Mohali",
        state="Punjab",
        source_permission_status="permitted_storage"
    )
    db_session.add(b_wa)
    db_session.flush()
    db_session.add(BusinessContact(
        branch_id=b_wa.id,
        contact_type="whatsapp",
        original_value="+91 98765 22222",
        normalized_value="+919876522222"
    ))

    # 4. Email-only lead (Strictly Excluded)
    b_em = Branch(
        company_id=comp.id,
        branch_name="Branch Email Only Excluded",
        normalized_branch_name="branch email only excluded",
        city="Mohali",
        state="Punjab",
        source_permission_status="permitted_storage"
    )
    db_session.add(b_em)
    db_session.flush()
    db_session.add(BusinessContact(
        branch_id=b_em.id,
        contact_type="email",
        original_value="excluded@hospital.com",
        normalized_value="excluded@hospital.com"
    ))

    # 5. No contact lead (Strictly Excluded)
    b_none = Branch(
        company_id=comp.id,
        branch_name="Branch No Contact Excluded",
        normalized_branch_name="branch no contact excluded",
        city="Mohali",
        state="Punjab",
        source_permission_status="permitted_storage"
    )
    db_session.add(b_none)

    # 6. Category with 0 contact numbers (For Rule 13 empty file test)
    comp_empty = Company(
        organization_id=org.id,
        legal_name="Empty Category Corp",
        normalized_name="empty category corp",
        primary_category="EmptyCategory"
    )
    db_session.add(comp_empty)
    db_session.flush()
    b_empty_cat = Branch(
        company_id=comp_empty.id,
        branch_name="Empty Category Lead",
        normalized_branch_name="empty category lead",
        city="Mohali",
        state="Punjab",
        source_permission_status="permitted_storage"
    )
    db_session.add(b_empty_cat)
    db_session.flush()
    # Give it only email, no phone
    db_session.add(BusinessContact(
        branch_id=b_empty_cat.id,
        contact_type="email",
        original_value="empty@empty.com",
        normalized_value="empty@empty.com"
    ))

    db_session.commit()

    # Test A: Combined Export must contain Mobile, Landline, WhatsApp, and EXCLUDE Email-only & No-contact
    res_comb = client.post(
        "/api/v1/exports",
        json={"export_mode": "all_combined", "file_format": "csv"}
    )
    assert res_comb.status_code == 201
    comb_info = res_comb.json()
    assert comb_info["status"] == "ready"

    dl_comb = client.get(comb_info["download_url"])
    csv_text = dl_comb.content.decode("utf-8-sig")

    assert "Branch Mobile Only" in csv_text
    assert "Branch Landline Only" in csv_text
    assert "Branch WhatsApp Only" in csv_text
    assert "Branch Email Only Excluded" not in csv_text
    assert "Branch No Contact Excluded" not in csv_text

    # Test B: Selected Leads Export passing ALL 5 leads
    # Only the 3 qualified ones should be exported!
    all_five_ids = [b_mob.id, b_land.id, b_wa.id, b_em.id, b_none.id]
    res_sel = client.post(
        "/api/v1/exports",
        json={"selected_ids": all_five_ids, "file_format": "xlsx"}
    )
    assert res_sel.status_code == 201
    sel_info = res_sel.json()
    assert sel_info["row_count"] == 3  # Exactly 3 out of 5 qualified!

    # Test C: Selected Leads with ONLY email and no-contact leads (Must fail gracefully, no empty file)
    unqual_ids = [b_em.id, b_none.id]
    res_unqual = client.post(
        "/api/v1/exports",
        json={"selected_ids": unqual_ids, "file_format": "csv"}
    )
    assert res_unqual.status_code == 201
    unqual_info = res_unqual.json()
    assert unqual_info["status"] == "failed"
    assert "None of the selected leads have a valid contact number" in unqual_info["error_message"]

    # Test D: Rule 13: Exporting a category with 0 qualifying contact numbers must fail gracefully, no empty file
    res_empty_cat = client.post(
        "/api/v1/exports",
        json={"export_mode": "single", "category": "EmptyCategory", "file_format": "xlsx"}
    )
    assert res_empty_cat.status_code == 201
    empty_cat_info = res_empty_cat.json()
    assert empty_cat_info["status"] == "failed"
    assert "0 contact-qualified leads with valid contact numbers" in empty_cat_info["error_message"]


def test_industry_exports_flow(client, db_session):
    # Setup test org & data
    org = db_session.query(Organization).first()
    if not org:
        org = Organization(id="org-test-export", name="Export Test Org", slug="export-test")
        db_session.add(org)
        db_session.flush()

    # 1. Create a Hospital with a mobile contact
    c_hosp = Company(
        organization_id=org.id,
        legal_name="Apollo Hospital",
        normalized_name="apollo hospital",
        primary_category="Hospital"
    )
    db_session.add(c_hosp)
    db_session.flush()
    b_hosp = Branch(
        company_id=c_hosp.id,
        branch_name="=cmd|'/C calc'!A0",  # Test formula injection in branch name!
        normalized_branch_name="apollo hospital mohali",
        city="Mohali",
        state="Punjab",
        source_permission_status="permitted_storage"
    )
    db_session.add(b_hosp)
    db_session.flush()
    db_session.add(BusinessContact(
        branch_id=b_hosp.id,
        contact_type="mobile",
        original_value="+91 98765 11111",
        normalized_value="+919876511111"
    ))

    # 2. Create a Salon with a landline contact
    c_salon = Company(
        organization_id=org.id,
        legal_name="Looks Salon",
        normalized_name="looks salon",
        primary_category="Salon"
    )
    db_session.add(c_salon)
    db_session.flush()
    b_salon = Branch(
        company_id=c_salon.id,
        branch_name="Looks Salon Sector 17",
        normalized_branch_name="looks salon sector 17",
        city="Chandigarh",
        state="Chandigarh",
        source_permission_status="permitted_storage"
    )
    db_session.add(b_salon)
    db_session.flush()
    db_session.add(BusinessContact(
        branch_id=b_salon.id,
        contact_type="landline",
        original_value="0172 2701111",
        normalized_value="01722701111"
    ))

    # 3. Create a Clinic with an advertised WhatsApp contact
    c_clinic = Company(
        organization_id=org.id,
        legal_name="City Clinic",
        normalized_name="city clinic",
        primary_category="Clinic"
    )
    db_session.add(c_clinic)
    db_session.flush()
    b_clinic = Branch(
        company_id=c_clinic.id,
        branch_name="City Clinic Phase 7",
        normalized_branch_name="city clinic phase 7",
        city="Mohali",
        state="Punjab",
        source_permission_status="permitted_storage"
    )
    db_session.add(b_clinic)
    db_session.flush()
    db_session.add(BusinessContact(
        branch_id=b_clinic.id,
        contact_type="whatsapp",
        original_value="+91 98765 22222",
        normalized_value="+919876522222"
    ))

    # 4. Create an Unqualified Hospital (Email only - no phone numbers)
    c_unqual = Company(
        organization_id=org.id,
        legal_name="Unqualified Hospital",
        normalized_name="unqualified hospital",
        primary_category="Hospital"
    )
    db_session.add(c_unqual)
    db_session.flush()
    b_unqual = Branch(
        company_id=c_unqual.id,
        branch_name="Unqualified Hospital Branch",
        normalized_branch_name="unqualified hospital branch",
        city="Mohali",
        state="Punjab",
        source_permission_status="permitted_storage"
    )
    db_session.add(b_unqual)
    db_session.flush()
    db_session.add(BusinessContact(
        branch_id=b_unqual.id,
        contact_type="email",
        original_value="contact@unqualified.com",
        normalized_value="contact@unqualified.com"
    ))

    # 5. Create an Unsaved Discovery Preview Item (Must NEVER be in exports)
    preview_item = DiscoveryJobItem(
        job_id="preview-job-1",
        business_name="Unsaved Discovery Preview Hospital",
        category="Hospital",
        phone="+919999999999",
        is_saved=False
    )
    db_session.add(preview_item)

    db_session.commit()

    # --- Test 1: GET /api/v1/exports/industries ---
    res_ind = client.get("/api/v1/exports/industries")
    assert res_ind.status_code == 200
    ind_data = res_ind.json()
    cats = {item["category"]: item for item in ind_data}
    assert "Hospital" in cats
    assert "Salon" in cats
    assert "Clinic" in cats
    # Verify qualified count reflects phone-qualified leads
    assert cats["Hospital"]["qualified_leads"] >= 1
    assert cats["Salon"]["qualified_leads"] >= 1
    assert cats["Clinic"]["qualified_leads"] >= 1

    # --- Test 2: Mode 'single' for Hospital in XLSX format -> hospital_leads.xlsx ---
    res_hosp = client.post(
        "/api/v1/exports",
        json={"export_mode": "single", "category": "Hospital", "file_format": "xlsx"}
    )
    assert res_hosp.status_code == 201
    hosp_info = res_hosp.json()
    assert hosp_info["status"] == "ready"
    assert hosp_info["filename"] == "hospital_leads.xlsx"

    # Download and inspect XLSX content
    dl_hosp = client.get(hosp_info["download_url"])
    assert dl_hosp.status_code == 200
    wb_hosp = openpyxl.load_workbook(io.BytesIO(dl_hosp.content))
    ws_leads = wb_hosp["India Business Leads"]
    rows = list(ws_leads.iter_rows(values_only=True))
    assert len(rows) >= 2  # header + data
    name_col_idx = rows[0].index("Business Name")
    cat_col_idx = rows[0].index("Category")
    for row in rows[1:]:
        assert row[cat_col_idx] == "Hospital", f"Expected only Hospital, got {row[cat_col_idx]}"
        # Ensure formula injection was sanitized in cell
        if "=cmd" in str(row[name_col_idx]):
            assert str(row[name_col_idx]).startswith("'=cmd")

    # Verify licensing sheet is present
    assert "License & Attribution" in wb_hosp.sheetnames

    # --- Test 3: Mode 'all_combined' in CSV format -> all_leads.csv ---
    res_comb = client.post(
        "/api/v1/exports",
        json={"export_mode": "all_combined", "file_format": "csv"}
    )
    assert res_comb.status_code == 201
    comb_info = res_comb.json()
    assert comb_info["status"] == "ready"
    assert comb_info["filename"] == "all_leads.csv"

    dl_comb = client.get(comb_info["download_url"])
    assert dl_comb.status_code == 200
    csv_text = dl_comb.content.decode("utf-8-sig")
    assert "Looks Salon Sector 17" in csv_text
    assert "City Clinic Phase 7" in csv_text
    assert "Unqualified Hospital Branch" not in csv_text  # Missing phone, must be excluded!
    assert "Unsaved Discovery Preview Hospital" not in csv_text  # Previews must NEVER be exported!

    # --- Test 4: Mode 'all_separate_zip' in XLSX format -> india_leads_by_industry.zip ---
    res_zip = client.post(
        "/api/v1/exports",
        json={"export_mode": "all_separate_zip", "file_format": "xlsx"}
    )
    assert res_zip.status_code == 201
    zip_info = res_zip.json()
    assert zip_info["status"] == "ready"
    assert zip_info["filename"] == "india_leads_by_industry.zip"
    assert zip_info["file_format"] == "zip"

    dl_zip = client.get(zip_info["download_url"])
    assert dl_zip.status_code == 200
    with zipfile.ZipFile(io.BytesIO(dl_zip.content), "r") as zf:
        namelist = zf.namelist()
        assert "hospital_leads.xlsx" in namelist
        assert "salon_leads.xlsx" in namelist
        assert "clinic_leads.xlsx" in namelist
        assert "README_LICENSING.txt" in namelist

        # Inspect README_LICENSING.txt
        readme_content = zf.read("README_LICENSING.txt").decode("utf-8")
        assert "Open Database License (ODbL) 1.0" in readme_content
        assert "OpenStreetMap" in readme_content

        # Inspect hospital_leads.xlsx inside ZIP
        zip_hosp_bytes = zf.read("hospital_leads.xlsx")
        wb_zip_hosp = openpyxl.load_workbook(io.BytesIO(zip_hosp_bytes))
        ws_zh = wb_zip_hosp["India Business Leads"]
        zh_rows = list(ws_zh.iter_rows(values_only=True))
        assert len(zh_rows) >= 2
        for row in zh_rows[1:]:
            assert row[cat_col_idx] == "Hospital"

    # --- Test 5: Mode 'all_separate_zip' in CSV format -> india_leads_by_industry.zip ---
    res_csv_zip = client.post(
        "/api/v1/exports",
        json={"export_mode": "all_separate_zip", "file_format": "csv"}
    )
    assert res_csv_zip.status_code == 201
    csv_zip_info = res_csv_zip.json()
    assert csv_zip_info["status"] == "ready"
    assert csv_zip_info["filename"] == "india_leads_by_industry.zip"

    dl_csv_zip = client.get(csv_zip_info["download_url"])
    assert dl_csv_zip.status_code == 200
    with zipfile.ZipFile(io.BytesIO(dl_csv_zip.content), "r") as zf:
        namelist = zf.namelist()
        assert "hospital_leads.csv" in namelist
        assert "salon_leads.csv" in namelist
        assert "clinic_leads.csv" in namelist
        assert "README_LICENSING.txt" in namelist
