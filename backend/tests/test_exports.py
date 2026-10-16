import pytest
import io
import openpyxl
from app.services.exports.csv_export import CsvExportService, sanitize_formula_injection
from app.services.exports.excel_export import ExcelExportService
from app.models import Branch, Company, BusinessContact


def test_sanitize_formula_injection():
    # Attack payloads attempting spreadsheet command execution
    payloads = [
        "=cmd|'/C calc'!A0",
        "+1+1",
        "-123",
        "@SUM(A1:A10)",
        "\t=2+2",
    ]

    for p in payloads:
        sanitized = sanitize_formula_injection(p)
        assert sanitized.startswith("'"), f"Failed to sanitize formula injection: {p} -> {sanitized}"

    # Normal strings left as-is
    assert sanitize_formula_injection("Glamour Salon") == "Glamour Salon"
    assert sanitize_formula_injection("+91 98765 43210").startswith("'")  # + is formula operator in Excel, escaped!


def test_csv_export_generation_and_utf8_bom(db_session):
    company = Company(
        organization_id="test-org-1",
        legal_name="Test Company",
        normalized_name="test company",
        primary_category="Salon"
    )
    db_session.add(company)
    db_session.flush()

    branch = Branch(
        company_id=company.id,
        branch_name="Test Branch Mohali",
        normalized_branch_name="test branch mohali",
        city="Mohali",
        state="Punjab",
        pincode="160062",
        source_permission_status="permitted_storage"
    )
    db_session.add(branch)
    db_session.flush()

    contact = BusinessContact(
        branch_id=branch.id,
        contact_type="mobile",
        original_value="+91 98765 43210",
        normalized_value="+919876543210",
        is_primary=True,
        source="test"
    )
    db_session.add(contact)
    db_session.commit()

    csv_bytes = CsvExportService.generate_csv_bytes([branch])
    # Must start with UTF-8 BOM \xef\xbb\xbf
    assert csv_bytes.startswith(b'\xef\xbb\xbf')

    csv_text = csv_bytes.decode("utf-8-sig")
    assert "Business Name" in csv_text
    assert "Test Branch Mohali" in csv_text
    assert "Mohali" in csv_text


def test_xlsx_export_generation(db_session):
    company = Company(
        organization_id="test-org-1",
        legal_name="Excel Test Corp",
        normalized_name="excel test corp",
        primary_category="Restaurant"
    )
    db_session.add(company)
    db_session.flush()

    branch = Branch(
        company_id=company.id,
        branch_name="Spice Kitchen Mohali",
        normalized_branch_name="spice kitchen mohali",
        city="Mohali",
        state="Punjab",
        source_permission_status="permitted_storage"
    )
    db_session.add(branch)
    db_session.commit()

    xlsx_bytes = ExcelExportService.generate_xlsx_bytes([branch])
    assert len(xlsx_bytes) > 1000

    # Verify openpyxl can read it back
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    assert ws.title == "India Business Leads"
    # Row 1 is header
    assert ws.cell(row=1, column=1).value == "Business Name"
    # Row 2 is data
    assert ws.cell(row=2, column=1).value == "Spice Kitchen Mohali"

